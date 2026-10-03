"""
Generate Caterpillar-dealer master data for the Supply Chain Decision Intelligence Platform.

Outputs (under 03_Source_Data/master_data/):
    golden/  - clean, governed reference data (the "truth" after stewardship)
    raw/     - messy extracts as they would arrive from ERP / spreadsheets
    dq_issue_log.csv - answer key: every defect injected into raw/, tagged with
                       the data quality dimension from 2.2 Data_Quality_Strategy.md

Equipment models and part-number formats follow real Caterpillar conventions
(e.g. CAT 320 excavator, 1R-0750 fuel filter). Costs, lead times, quantities and
supplier performance values are synthetic.

Usage:
    python generate_master_data.py            # default seed 42
    python generate_master_data.py --seed 7
    python generate_master_data.py --scale 3   # ~3x equipment and parts
"""

import argparse
import random
from datetime import date, timedelta
from pathlib import Path

import pandas as pd

OUT_DIR = Path(__file__).parent / "master_data"

# ---------------------------------------------------------------------------
# Reference content
# ---------------------------------------------------------------------------

# Dealer branch network (VA / WV / MD territory, matching the spec examples)
BRANCHES = [
    ("Abingdon", "Southwest", "VA"), ("Salem", "Southwest", "VA"),
    ("Lynchburg", "Central", "VA"), ("Richmond", "East", "VA"),
    ("Ashland", "East", "VA"), ("Chesapeake", "East", "VA"),
    ("Fredericksburg", "North", "VA"), ("Manassas", "North", "VA"),
    ("Harrisonburg", "North", "VA"), ("Winchester", "North", "VA"),
    ("Bluefield", "Southwest", "WV"), ("Beckley", "Southwest", "WV"),
    ("Hagerstown", "North", "MD"), ("Salisbury", "East", "MD"),
    ("Danville", "Central", "VA"), ("Roanoke", "Southwest", "VA"),
    ("Norfolk", "East", "VA"), ("Charlottesville", "Central", "VA"),
    ("Culpeper", "North", "VA"), ("South Boston", "Central", "VA"),
    ("Wytheville", "Southwest", "VA"), ("Grundy", "Southwest", "VA"),
    ("Princeton", "Southwest", "WV"), ("Logan", "Southwest", "WV"),
    ("Baltimore", "North", "MD"), ("Waldorf", "East", "MD"),
    ("Frederick", "North", "MD"), ("Dover", "East", "DE"),
]
BRANCH_TYPES = ["Full Service", "Full Service", "Parts & Service", "Rental Store"]

# (family name, family code, Caterpillar business segment, planning owner role)
PRODUCT_FAMILIES = [
    ("Excavation", "EXC", "Construction Industries", "Construction Planning Manager"),
    ("Earthmoving", "EMV", "Construction Industries", "Construction Planning Manager"),
    ("Material Handling", "MHD", "Construction Industries", "Compact Equipment Planner"),
    ("Power Systems", "PWR", "Energy & Transportation", "Power Systems Planner"),
    ("Road Construction", "RDC", "Construction Industries", "Paving Products Planner"),
    ("Mining", "MIN", "Resource Industries", "Mining Planning Manager"),
    ("Underground Mining", "UGM", "Resource Industries", "Mining Planning Manager"),
    ("Forestry", "FOR", "Construction Industries", "Forestry Planner"),
    ("Work Tools & Attachments", "WTA", "Construction Industries", "Work Tools Planner"),
    ("Marine & Oil and Gas", "MOG", "Energy & Transportation", "Power Systems Planner"),
]

# category -> (family, criticality, [models])
CATEGORIES = {
    "Mini Excavators":        ("Excavation", "Medium", ["301.7", "303 CR", "305 CR", "308 CR"]),
    "Excavators":             ("Excavation", "High",   ["313", "315", "320", "323", "326", "330", "336", "352", "395"]),
    "Wheeled Excavators":     ("Excavation", "Medium", ["M318", "M320"]),
    "Backhoe Loaders":        ("Excavation", "Medium", ["416", "420", "430", "440"]),
    "Dozers":                 ("Earthmoving", "High",  ["D3", "D4", "D5", "D6", "D6 XE", "D8", "D9", "D10", "D11"]),
    "Wheel Loaders":          ("Earthmoving", "High",  ["906", "908", "926", "938", "950", "966", "972", "980", "988"]),
    "Motor Graders":          ("Earthmoving", "High",  ["120", "140", "150", "160", "18"]),
    "Articulated Trucks":     ("Earthmoving", "High",  ["725", "730", "745"]),
    "Off-Highway Trucks":     ("Earthmoving", "High",  ["770", "772", "775", "777"]),
    "Track Loaders":          ("Earthmoving", "Medium", ["953", "963", "973"]),
    "Skid Steer Loaders":     ("Material Handling", "Low", ["226", "236", "246", "262", "272"]),
    "Compact Track Loaders":  ("Material Handling", "Medium", ["239", "259", "279", "289", "299"]),
    "Telehandlers":           ("Material Handling", "Low", ["TH255C", "TH357D", "TL943", "TL1055", "TL1255"]),
    "Generator Sets":         ("Power Systems", "High", ["C9.3B", "C15", "C18", "C32", "3516"]),
    "Industrial Engines":     ("Power Systems", "Medium", ["C4.4", "C7.1", "C13", "C18 ACERT"]),
    "Asphalt Pavers":         ("Road Construction", "Medium", ["AP400", "AP555", "AP1055"]),
    "Cold Planers":           ("Road Construction", "Medium", ["PM310", "PM620", "PM822"]),
    "Compactors":             ("Road Construction", "Low", ["CB2.7", "CB10", "CS56B", "CS66B", "CW34"]),
    "Road Reclaimers":        ("Road Construction", "Medium", ["RM400", "RM600"]),
    "Pipelayers":             ("Earthmoving", "Medium", ["PL61", "PL72", "PL87"]),
    "Material Handlers":      ("Material Handling", "Medium", ["MH3024", "MH3026", "MH3040"]),
    "Gas Generator Sets":     ("Power Systems", "High", ["G3516", "G3520"]),
    "Large Mining Trucks":    ("Mining", "High", ["793", "794 AC", "795", "797F"]),
    "Hydraulic Mining Shovels": ("Mining", "High", ["6015B", "6020B", "6030", "6040", "6060"]),
    "Electric Rope Shovels":  ("Mining", "High", ["7495"]),
    "Rotary Drills":          ("Mining", "High", ["MD6250", "MD6310"]),
    "Underground Loaders":    ("Underground Mining", "High", ["R1300G", "R1700", "R2900"]),
    "Underground Trucks":     ("Underground Mining", "High", ["AD22", "AD45", "AD60"]),
    "Wheel Skidders":         ("Forestry", "Medium", ["525D", "535D", "545D"]),
    "Track Feller Bunchers":  ("Forestry", "Medium", ["521B", "522B"]),
    "Forest Machines":        ("Forestry", "Medium", ["558", "568"]),
    "Hammers":                ("Work Tools & Attachments", "Low", ["H95", "H110", "H120", "H140"]),
    "Augers & Grapples":      ("Work Tools & Attachments", "Low", ["A14B Auger", "G115B Grapple", "G120B Grapple"]),
    "Marine Engines":         ("Marine & Oil and Gas", "High", ["C18 Marine", "C32 Marine", "3512E Marine"]),
    "Well Service Engines":   ("Marine & Oil and Gas", "High", ["3512E Well Service", "3516E Well Service"]),
}

# Parts: (part number, description, part category, criticality, base unit cost, uom)
# Mix of well-known Cat service part numbers and Cat-format generated numbers.
SEED_PARTS = [
    ("1R-0750", "Fuel Filter", "Filters", "High", 38.0, "EA"),
    ("1R-0751", "Fuel Filter Secondary", "Filters", "High", 41.0, "EA"),
    ("1R-0739", "Engine Oil Filter", "Filters", "High", 29.0, "EA"),
    ("1R-1808", "Engine Oil Filter", "Filters", "High", 33.0, "EA"),
    ("1R-0716", "Engine Oil Filter", "Filters", "High", 31.0, "EA"),
    ("326-1644", "Fuel Water Separator", "Filters", "High", 64.0, "EA"),
    ("6I-2501", "Primary Air Filter", "Filters", "Medium", 112.0, "EA"),
    ("6I-2502", "Secondary Air Filter", "Filters", "Medium", 74.0, "EA"),
    ("131-8822", "Air Filter Element", "Filters", "Medium", 98.0, "EA"),
    ("1R-0777", "Hydraulic Filter", "Filters", "High", 86.0, "EA"),
    ("1U-3352", "Bucket Tip J350", "Ground Engaging Tools", "Medium", 46.0, "EA"),
    ("1U-3302", "Bucket Tip J300", "Ground Engaging Tools", "Medium", 34.0, "EA"),
    ("8E-6208", "Bolt Plow", "Hardware", "Low", 3.5, "EA"),
    ("4K-0367", "Nut Plow", "Hardware", "Low", 1.2, "EA"),
]
GENERATED_PART_TYPES = [
    ("Track Shoe", "Undercarriage", "High", 180), ("Track Roller", "Undercarriage", "High", 420),
    ("Idler Group", "Undercarriage", "High", 1900), ("Sprocket Segment", "Undercarriage", "High", 260),
    ("Track Link Assembly", "Undercarriage", "High", 5200), ("Cutting Edge", "Ground Engaging Tools", "Medium", 310),
    ("End Bit", "Ground Engaging Tools", "Medium", 140), ("Adapter", "Ground Engaging Tools", "Medium", 95),
    ("Hydraulic Hose Assembly", "Hydraulics", "High", 120), ("Hydraulic Pump", "Hydraulics", "High", 6800),
    ("Seal Kit", "Hydraulics", "Medium", 75), ("Hydraulic Cylinder", "Hydraulics", "High", 3400),
    ("Fuel Injector", "Engine", "High", 640), ("Water Pump", "Engine", "High", 520),
    ("Turbocharger", "Engine", "High", 2900), ("Gasket Kit", "Engine", "Medium", 160),
    ("V-Belt", "Engine", "Low", 28), ("Alternator", "Electrical", "Medium", 610),
    ("Starter Motor", "Electrical", "High", 780), ("Wiring Harness", "Electrical", "Medium", 450),
    ("Sensor Speed", "Electrical", "Medium", 140), ("Battery", "Electrical", "Medium", 210),
    ("Cab Air Filter", "Filters", "Low", 42), ("Transmission Filter", "Filters", "Medium", 58),
    ("Radiator", "Cooling", "High", 2400), ("Thermostat", "Cooling", "Medium", 48),
    ("Track Pin", "Undercarriage", "Medium", 65), ("Track Bushing", "Undercarriage", "Medium", 55),
    ("Carrier Roller", "Undercarriage", "High", 380), ("Recoil Spring", "Undercarriage", "High", 1500),
    ("Ripper Shank", "Ground Engaging Tools", "Medium", 900), ("Ripper Tip", "Ground Engaging Tools", "Medium", 60),
    ("Bucket Tooth Retainer", "Ground Engaging Tools", "Low", 6), ("Side Cutter", "Ground Engaging Tools", "Medium", 210),
    ("Control Valve", "Hydraulics", "High", 4200), ("Swing Motor", "Hydraulics", "High", 7600),
    ("Final Drive", "Power Train", "High", 14500), ("Travel Motor", "Power Train", "High", 9800),
    ("Torque Converter", "Power Train", "High", 8200), ("Transmission Clutch Disc", "Power Train", "Medium", 140),
    ("Axle Seal", "Power Train", "Medium", 85), ("Brake Disc", "Power Train", "Medium", 320),
    ("Cylinder Head", "Engine", "High", 3900), ("Piston Kit", "Engine", "High", 780),
    ("Crankshaft Bearing", "Engine", "Medium", 95), ("Fuel Transfer Pump", "Engine", "High", 450),
    ("ECM Controller", "Electrical", "High", 2600), ("Monitor Display", "Electrical", "Medium", 1300),
    ("Relay", "Electrical", "Low", 22), ("Work Light LED", "Electrical", "Low", 180),
    ("Fan Blade", "Cooling", "Medium", 260), ("Coolant Hose", "Cooling", "Low", 45),
    ("Seat Assembly", "Cab", "Low", 1100), ("Window Glass", "Cab", "Low", 420),
    ("Mirror", "Cab", "Low", 75), ("O-Ring Seal", "Hardware", "Low", 2),
    ("Hex Bolt", "Hardware", "Low", 4), ("Washer Hard", "Hardware", "Low", 1),
    ("Tire 29.5R25", "Tires", "High", 9500), ("Tire 12-16.5", "Tires", "Medium", 320),
]
FLUIDS = [
    ("DEO 15W-40 Engine Oil", "Fluids", "Medium", 24.0, "GAL"),
    ("HYDO Advanced 10 Hydraulic Oil", "Fluids", "Medium", 27.0, "GAL"),
    ("ELC Extended Life Coolant", "Fluids", "Medium", 19.0, "GAL"),
    ("TDTO Transmission Oil", "Fluids", "Medium", 26.0, "GAL"),
]

# (name, tier, supplier type, city, state, base lead time days)
SUPPLIERS = [
    ("Caterpillar Inc. - Parts Distribution Center Morton", "Tier 1", "OEM", "Morton", "IL", 3),
    ("Caterpillar Inc. - Parts Distribution Center York", "Tier 1", "OEM", "York", "PA", 2),
    ("Caterpillar Inc. - Parts Distribution Center Atlanta", "Tier 1", "OEM", "Atlanta", "GA", 3),
    ("Caterpillar Inc. - Machine Order Fulfillment", "Tier 1", "OEM", "Peoria", "IL", 45),
    ("Caterpillar Reman Services", "Tier 1", "OEM Reman", "Corinth", "MS", 14),
    ("Donaldson Company", "Tier 2", "Filtration", "Bloomington", "MN", 10),
    ("Parker Hannifin", "Tier 2", "Hydraulics", "Cleveland", "OH", 12),
    ("Gates Corporation", "Tier 2", "Hoses & Belts", "Denver", "CO", 9),
    ("Berco of America", "Tier 2", "Undercarriage", "Mount Prospect", "IL", 28),
    ("ITR America", "Tier 2", "Undercarriage", "Lombard", "IL", 21),
    ("ESCO Group", "Tier 2", "Ground Engaging Tools", "Portland", "OR", 18),
    ("Midwest Components", "Tier 3", "Distributor", "Peoria", "IL", 10),
    ("Atlantic Hydraulic Supply", "Tier 3", "Distributor", "Norfolk", "VA", 5),
    ("Blue Ridge Industrial Supply", "Tier 3", "Distributor", "Roanoke", "VA", 4),
    ("Interstate Batteries", "Tier 3", "Electrical", "Dallas", "TX", 6),
    ("Appalachian Fastener Co.", "Tier 3", "Hardware", "Bristol", "TN", 7),
    ("Caterpillar Inc. - Parts Distribution Center Denver", "Tier 1", "OEM", "Denver", "CO", 4),
    ("Caterpillar Inc. - Parts Distribution Center Waco", "Tier 1", "OEM", "Waco", "TX", 4),
    ("Perkins Engines", "Tier 1", "OEM Engines", "Peterborough", "UK", 35),
    ("Solar Turbines", "Tier 1", "OEM Power", "San Diego", "CA", 60),
    ("Cummins Filtration", "Tier 2", "Filtration", "Nashville", "TN", 9),
    ("Eaton Hydraulics", "Tier 2", "Hydraulics", "Eden Prairie", "MN", 14),
    ("Bosch Rexroth", "Tier 2", "Hydraulics", "Bethlehem", "PA", 20),
    ("Dana Off-Highway", "Tier 2", "Power Train", "Maumee", "OH", 30),
    ("ZF North America", "Tier 2", "Power Train", "Northville", "MI", 32),
    ("Hensley Industries", "Tier 2", "Ground Engaging Tools", "Dallas", "TX", 15),
    ("Bridgestone OTR", "Tier 2", "Tires", "Nashville", "TN", 40),
    ("Michelin Earthmover", "Tier 2", "Tires", "Greenville", "SC", 45),
    ("Delco Remy", "Tier 2", "Electrical", "Pendleton", "IN", 12),
    ("Grammer Seating", "Tier 3", "Cab", "Shelbyville", "TN", 18),
    ("Old Dominion Hose & Fitting", "Tier 3", "Hydraulics", "Richmond", "VA", 3),
    ("Tidewater Industrial Parts", "Tier 3", "Distributor", "Chesapeake", "VA", 4),
    ("Mountain State Equipment Supply", "Tier 3", "Distributor", "Beckley", "WV", 5),
    ("Chesapeake Bay Electrical", "Tier 3", "Electrical", "Baltimore", "MD", 6),
    ("Shenandoah Cooling Systems", "Tier 3", "Cooling", "Harrisonburg", "VA", 8),
    ("Keystone Glass & Cab Parts", "Tier 3", "Cab", "Harrisburg", "PA", 10),
]

SERIAL_PREFIX_CHARS = "ABCDEFGHJKLMNPRSTVWXYZ"


# ---------------------------------------------------------------------------
# Golden (clean) master data
# ---------------------------------------------------------------------------

def build_golden(rng: random.Random, scale: float = 1.0) -> dict[str, pd.DataFrame]:
    branch = pd.DataFrame([
        {
            "branch_id": f"B{i:03d}",
            "branch_name": name,
            "region": region,
            "state": state,
            "branch_type": rng.choice(BRANCH_TYPES),
            "open_date": date(rng.randint(1958, 2018), rng.randint(1, 12), 1).isoformat(),
            "is_active": "Y",
        }
        for i, (name, region, state) in enumerate(BRANCHES, start=1)
    ])

    family = pd.DataFrame([
        {"product_family_id": f"PF{i:03d}", "product_family_name": f, "product_family_code": code,
         "business_segment": seg, "planning_owner": owner, "is_active": "Y"}
        for i, (f, code, seg, owner) in enumerate(PRODUCT_FAMILIES, start=1)
    ])
    fam_id = dict(zip(family.product_family_name, family.product_family_id))

    category = pd.DataFrame([
        {"product_category_id": f"PC{i:03d}", "product_category_name": c,
         "product_family_id": fam_id[fam]}
        for i, (c, (fam, _, _)) in enumerate(CATEGORIES.items(), start=1)
    ])
    cat_id = dict(zip(category.product_category_name, category.product_category_id))

    product_rows, equipment_rows = [], []
    pid = 1
    for cat_name, (fam, crit, models) in CATEGORIES.items():
        for m in models:
            product_rows.append({
                "product_id": f"P{pid:03d}",
                "product_name": f"CAT {m}",
                "product_category_id": cat_id[cat_name],
                "product_family_id": fam_id[fam],
                "criticality": crit,
                "lifecycle_status": rng.choices(["Current", "Current", "Current", "Legacy"], k=1)[0],
            })
            pid += 1
    product = pd.DataFrame(product_rows)

    # Equipment master: individual serialized units in the field, tied to a branch
    eid = 1
    for _, p in product.iterrows():
        prefix = "".join(rng.choice(SERIAL_PREFIX_CHARS) for _ in range(3))
        for _ in range(max(1, int(rng.randint(15, 60) * scale))):
            year = rng.randint(2008, 2026)
            equipment_rows.append({
                "equipment_id": f"E{eid:05d}",
                "serial_number": f"{prefix}{rng.randint(1, 99999):05d}",
                "equipment_model": p.product_name,
                "product_id": p.product_id,
                "product_category_id": p.product_category_id,
                "model_year": year,
                "service_meter_hours": max(0, int(rng.gauss((2026 - year) * 1100, 600))),
                "home_branch_id": rng.choice(branch.branch_id.tolist()),
                "ownership": rng.choices(["Customer", "Rental Fleet", "Dealer Demo"], [80, 17, 3])[0],
            })
            eid += 1
    equipment = pd.DataFrame(equipment_rows)

    supplier = pd.DataFrame([
        {
            "supplier_id": f"S{i:03d}",
            "supplier_name": name,
            "supplier_tier": tier,
            "supplier_type": stype,
            "city": city,
            "state": st,
            "lead_time_days": lt,
            "on_time_delivery_pct": round(min(99.5, rng.gauss(94 if tier == "Tier 1" else 86, 4)), 1),
            "payment_terms": rng.choice(["Net 30", "Net 45", "Net 60"]),
            "is_active": "Y",
        }
        for i, (name, tier, stype, city, st, lt) in enumerate(SUPPLIERS, start=1)
    ])
    sup_by_type = supplier.groupby("supplier_type").supplier_id.apply(list).to_dict()
    cat_pdcs = supplier[supplier.supplier_type == "OEM"].supplier_id.tolist()[:3]

    def pick_supplier(part_cat: str) -> str:
        alt = {"Filters": "Filtration", "Hydraulics": "Hydraulics", "Undercarriage": "Undercarriage",
               "Ground Engaging Tools": "Ground Engaging Tools", "Electrical": "Electrical",
               "Hardware": "Hardware", "Power Train": "Power Train", "Tires": "Tires",
               "Cooling": "Cooling", "Cab": "Cab"}.get(part_cat)
        if alt and alt in sup_by_type and rng.random() < 0.3:
            return rng.choice(sup_by_type[alt])
        return rng.choice(cat_pdcs)  # most genuine Cat parts ship from a Cat PDC

    part_rows, used = [], set(p[0] for p in SEED_PARTS)

    def new_part_number() -> str:
        while True:
            if rng.random() < 0.5:
                pn = f"{rng.randint(100, 599)}-{rng.randint(0, 9999):04d}"
            else:
                pn = f"{rng.randint(1, 9)}{rng.choice('ABCDEFGHIJKLMNPRSTUVWXY')}-{rng.randint(0, 9999):04d}"
            if pn not in used:
                used.add(pn)
                return pn

    all_models = product.product_id.tolist()
    for pn, desc, pcat, crit, cost, uom in SEED_PARTS:
        part_rows.append((pn, desc, pcat, crit, cost, uom))
    for desc, pcat, crit, cost in GENERATED_PART_TYPES:
        for _ in range(max(1, int(rng.randint(40, 120) * scale))):
            part_rows.append((new_part_number(), desc, pcat, crit,
                              round(cost * rng.uniform(0.6, 1.8), 2), "EA"))
    for desc, pcat, crit, cost, uom in FLUIDS:
        part_rows.append((new_part_number(), desc, pcat, crit, cost, uom))

    part = pd.DataFrame([
        {
            "part_id": f"PT{i:05d}",
            "part_number": pn,
            "part_description": desc.upper(),
            "part_category": pcat,
            "criticality": crit,
            "unit_cost_usd": cost,
            "unit_of_measure": uom,
            "primary_supplier_id": pick_supplier(pcat),
            "primary_product_id": rng.choice(all_models),
            "reman_available": "Y" if pcat in ("Engine", "Hydraulics", "Electrical") and rng.random() < 0.4 else "N",
            "created_date": (date(2015, 1, 1) + timedelta(days=rng.randint(0, 3800))).isoformat(),
            "is_active": "Y",
        }
        for i, (pn, desc, pcat, crit, cost, uom) in enumerate(part_rows, start=1)
    ])

    return {
        "branch_master": branch,
        "product_family_master": family,
        "product_category_master": category,
        "product_master": product,
        "equipment_master": equipment,
        "supplier_master": supplier,
        "part_master": part,
    }


# ---------------------------------------------------------------------------
# Mess injection
# ---------------------------------------------------------------------------

class Messer:
    """Applies realistic defects to a copy of the golden data and logs each one."""

    def __init__(self, rng: random.Random):
        self.rng = rng
        self.log: list[dict] = []

    def record(self, table, key, column, issue, dimension, severity, original, injected):
        self.log.append({
            "table": table, "record_key": key, "column": column, "issue_type": issue,
            "dq_dimension": dimension, "severity": severity,
            "golden_value": original, "raw_value": injected,
        })

    def sample(self, df, frac, minimum=1, col=None, kind=None):
        """Pick rows to damage; with col, only rows whose value is still usable (non-null, of kind)."""
        idx = list(df.index)
        if col:
            idx = [i for i in idx if df.at[i, col] is not None and not pd.isna(df.at[i, col])
                   and (kind is None or isinstance(df.at[i, col], kind))]
        n = max(minimum, int(len(df) * frac))
        return self.rng.sample(idx, min(n, len(idx)))

    def set(self, df, idx, col, value, table, key_col, issue, dim, sev):
        original = df.at[idx, col]
        df.at[idx, col] = value
        self.record(table, df.at[idx, key_col], col, issue, dim, sev, original, value)

    def duplicate(self, df, idx, table, key_col, mutate, issue, sev):
        row = df.loc[idx].copy()
        original_key = row[key_col]
        for col, fn in mutate.items():
            row[col] = fn(row[col])
        self.record(table, row[key_col], "*", issue, "Uniqueness", sev, original_key, row[key_col])
        return row


STATE_NAMES = {"VA": "Virginia", "WV": "West Virginia", "MD": "Maryland", "IL": "Illinois",
               "PA": "Pennsylvania", "GA": "Georgia", "MS": "Mississippi", "MN": "Minnesota",
               "OH": "Ohio", "CO": "Colorado", "OR": "Oregon", "TX": "Texas", "TN": "Tennessee",
               "DE": "Delaware", "CA": "California", "SC": "South Carolina", "MI": "Michigan",
               "IN": "Indiana", "UK": "United Kingdom"}


def messy_date(iso: str, rng: random.Random) -> str:
    d = date.fromisoformat(iso)
    return rng.choice([d.strftime("%m/%d/%Y"), d.strftime("%d-%b-%y"), d.strftime("%Y%m%d"),
                       d.strftime("%m/%d/%y")])


def supplier_alias(name: str, rng: random.Random) -> str:
    if name.startswith("Caterpillar Inc. - Parts Distribution Center"):
        city = name.split()[-1]
        return rng.choice([f"CAT PDC {city.upper()}", f"Caterpillar {city} PDC",
                           f"CATERPILLAR INC {city.upper()}", f"Cat Parts Dist Ctr - {city}"])
    base = name.replace(" Company", "").replace(" Corporation", "").replace(" Co.", "")
    return rng.choice([base.upper(), f"{base} Inc", f"{base} LLC", f"{base}, Inc.", base + " "])


def build_raw(golden: dict[str, pd.DataFrame], rng: random.Random):
    m = Messer(rng)
    raw = {k: v.copy().astype(object) for k, v in golden.items()}

    # --- Branch -----------------------------------------------------------
    t, df = "branch_master", raw["branch_master"]
    for i in m.sample(df, 0.25):
        m.set(df, i, "state", STATE_NAMES[df.at[i, "state"]], t, "branch_id",
              "State spelled out instead of code", "Consistency", "Low")
    for i in m.sample(df, 0.2):
        m.set(df, i, "branch_name", rng.choice([f"  {df.at[i, 'branch_name']}", df.at[i, "branch_name"].upper(),
                                                 df.at[i, "branch_name"] + " Branch"]),
              t, "branch_id", "Branch name formatting", "Consistency", "Low")
    i = m.sample(df, 0.07)[0]
    m.set(df, i, "region", None, t, "branch_id", "Missing region", "Completeness", "High")
    i = m.sample(df, 0.07)[0]
    m.set(df, i, "region", "Eastern", t, "branch_id", "Region value not in allowed list", "Validity", "Medium")
    for i in df.index:
        df.at[i, "open_date"] = messy_date(df.at[i, "open_date"], rng)
    m.record(t, "*", "open_date", "Mixed date formats across extract", "Validity", "Low", "YYYY-MM-DD", "mixed")
    dup = m.duplicate(df, m.sample(df, 0.07)[0], t, "branch_id",
                      {"branch_id": lambda x: "B9" + x[2:], "branch_name": lambda x: x.strip().title() + " Store"},
                      "Duplicate branch under new ID", "High")
    df.loc[len(df)] = dup
    df.loc[len(df)] = {"branch_id": "B050", "branch_name": "Norton", "region": "Southwest", "state": "VA",
                       "branch_type": "Parts & Service", "open_date": "2004-06-01", "is_active": "N"}
    m.record(t, "B050", "is_active", "Closed branch still present in master", "Validity", "Low", None, "N")

    # --- Product family / category ---------------------------------------
    t, df = "product_family_master", raw["product_family_master"]
    fam_idx = {df.at[i, "product_family_name"]: i for i in df.index}
    # Duplicate families created by different teams / legacy ERP loads
    family_dups = [
        ("PF011", "Earthmoving", {"product_family_name": "EARTHMOVING", "product_family_code": "EMV"},
         "Duplicate family with different casing"),
        ("PF012", "Excavation", {"product_family_name": "Excavation ", "product_family_code": "EXCV"},
         "Duplicate family with trailing space and alternate code"),
        ("PF013", "Road Construction", {"product_family_name": "Paving", "product_family_code": "PAV"},
         "Legacy family name (Paving) duplicating Road Construction"),
        ("PF014", "Power Systems", {"product_family_name": "Power Sys.", "product_family_code": None},
         "Abbreviated duplicate family missing code"),
    ]
    for new_id, base, changes, issue in family_dups:
        row = df.loc[fam_idx[base]].copy()
        row["product_family_id"] = new_id
        for col, val in changes.items():
            row[col] = val
        df.loc[len(df)] = row
        m.record(t, new_id, "product_family_name", issue, "Uniqueness", "High", base, row["product_family_name"])
    m.set(df, fam_idx["Earthmoving"], "product_family_name", "Earth Moving", t, "product_family_id",
          "Family name spelling differs from standard", "Consistency", "Medium")
    m.set(df, fam_idx["Mining"], "business_segment", "Construction Industries", t, "product_family_id",
          "Family mapped to wrong business segment", "Validity", "High")
    m.set(df, fam_idx["Forestry"], "business_segment", None, t, "product_family_id",
          "Missing business segment", "Completeness", "Medium")
    m.set(df, fam_idx["Material Handling"], "business_segment", "CI", t, "product_family_id",
          "Business segment abbreviated", "Consistency", "Low")
    m.set(df, fam_idx["Marine & Oil and Gas"], "product_family_name", "Marine and O&G", t, "product_family_id",
          "Non-standard family name", "Consistency", "Low")
    m.set(df, fam_idx["Work Tools & Attachments"], "planning_owner", None, t, "product_family_id",
          "Missing planning owner", "Completeness", "Medium")
    m.set(df, fam_idx["Underground Mining"], "is_active", "Yes", t, "product_family_id",
          "Active flag not Y/N", "Validity", "Low")
    df.loc[len(df)] = {"product_family_id": None, "product_family_name": "Used Equipment",
                       "product_family_code": "USED", "business_segment": None,
                       "planning_owner": None, "is_active": "Y"}
    m.record(t, None, "product_family_id", "Family row loaded without ID", "Completeness", "High", None, "Used Equipment")
    df.loc[len(df)] = {"product_family_id": "PF015", "product_family_name": "Agriculture",
                       "product_family_code": "AGR", "business_segment": "Construction Industries",
                       "planning_owner": "Construction Planning Manager", "is_active": "N"}
    m.record(t, "PF015", "is_active", "Retired family (Ag line exited) still in master", "Validity", "Low", None, "N")

    t, df = "product_category_master", raw["product_category_master"]
    cat_idx = {df.at[i, "product_category_name"]: i for i in df.index}
    m.set(df, cat_idx["Compact Track Loaders"], "product_category_name", "Compact Track Loader (CTL)", t,
          "product_category_id", "Non-standard category name", "Consistency", "Low")
    m.set(df, cat_idx["Off-Highway Trucks"], "product_category_name", "OHT", t,
          "product_category_id", "Category name abbreviated", "Consistency", "Low")
    m.set(df, cat_idx["Track Loaders"], "product_family_id", None, t, "product_category_id",
          "Category missing product family", "Completeness", "High")
    m.set(df, cat_idx["Telehandlers"], "product_family_id", "PF011", t, "product_category_id",
          "Category mapped to duplicate family", "Integrity", "Medium")
    m.set(df, cat_idx["Asphalt Pavers"], "product_family_id", "PF013", t, "product_category_id",
          "Category mapped to legacy family", "Integrity", "Medium")
    m.set(df, cat_idx["Large Mining Trucks"], "product_family_id", "PF002", t, "product_category_id",
          "Category mapped to wrong family (Earthmoving vs Mining)", "Validity", "High")
    m.set(df, cat_idx["Hammers"], "product_family_id", "PF099", t, "product_category_id",
          "Orphan family reference", "Integrity", "High")
    row = df.loc[cat_idx["Excavators"]].copy()
    row["product_category_id"], row["product_category_name"] = "PC099", "Hydraulic Excavators"
    df.loc[len(df)] = row
    m.record(t, "PC099", "product_category_name", "Duplicate category under synonym", "Uniqueness", "High",
             "Excavators", "Hydraulic Excavators")

    # --- Product ----------------------------------------------------------
    t, df = "product_master", raw["product_master"]
    for i in m.sample(df, 0.12):
        orig = df.at[i, "product_name"]
        m.set(df, i, "product_name", rng.choice([orig.replace("CAT ", "Cat"), orig.replace("CAT ", ""),
                                                  orig.replace("CAT ", "Caterpillar "), orig + "  "]),
              t, "product_id", "Model name not standardized", "Consistency", "Medium")
    for i in m.sample(df, 0.05):
        m.set(df, i, "criticality", None, t, "product_id", "Missing criticality", "Completeness", "High")
    for i in m.sample(df, 0.04):
        m.set(df, i, "criticality", rng.choice(["Hi", "CRITICAL", "med"]), t, "product_id",
              "Criticality not in allowed list", "Validity", "Medium")
    for i in m.sample(df, 0.03):
        m.set(df, i, "product_category_id", "PC199", t, "product_id",
              "Orphan category reference", "Integrity", "High")
    for i in m.sample(df, 0.03):
        dup = m.duplicate(df, i, t, "product_id",
                          {"product_id": lambda x: "P" + str(500 + int(x[1:]))}, "Duplicate product under new ID", "High")
        df.loc[len(df)] = dup

    # --- Supplier ---------------------------------------------------------
    t, df = "supplier_master", raw["supplier_master"]
    for i in m.sample(df, 0.3):
        dup = m.duplicate(df, i, t, "supplier_id",
                          {"supplier_id": lambda x: "S" + str(100 + int(x[1:])),
                           "supplier_name": lambda x: supplier_alias(x, rng)},
                          "Duplicate supplier with name variant", "High")
        df.loc[len(df)] = dup
    for i in m.sample(df, 0.12):
        m.set(df, i, "lead_time_days", None, t, "supplier_id", "Missing lead time", "Completeness", "High")
    for i in m.sample(df, 0.1):
        m.set(df, i, "lead_time_days", f"{df.at[i, 'lead_time_days']} days" if df.at[i, "lead_time_days"] is not None
              else "TBD", t, "supplier_id", "Lead time stored as text", "Validity", "Medium")
    i = m.sample(df, 0.05)[0]
    m.set(df, i, "lead_time_days", -7, t, "supplier_id", "Negative lead time", "Validity", "High")
    i = m.sample(df, 0.05)[0]
    m.set(df, i, "on_time_delivery_pct", 112.0, t, "supplier_id", "On-time % above 100", "Validity", "Medium")
    for i in m.sample(df, 0.15):
        m.set(df, i, "supplier_tier", rng.choice(["T1", "tier 2", "Tier-3", "1"]), t, "supplier_id",
              "Tier value not standardized", "Consistency", "Low")
    for i in m.sample(df, 0.15):
        m.set(df, i, "state", STATE_NAMES.get(df.at[i, "state"], df.at[i, "state"]), t, "supplier_id",
              "State spelled out instead of code", "Consistency", "Low")
    i = m.sample(df, 0.05)[0]
    m.set(df, i, "is_active", "N", t, "supplier_id", "Supplier inactive but still referenced by parts",
          "Integrity", "High")

    # --- Equipment --------------------------------------------------------
    t, df = "equipment_master", raw["equipment_master"]
    for i in m.sample(df, 0.04):
        m.set(df, i, "home_branch_id", rng.choice([None, "B099", "Richmond"]), t, "equipment_id",
              "Missing or invalid home branch", "Integrity", "High")
    for i in m.sample(df, 0.03):
        m.set(df, i, "serial_number", None, t, "equipment_id", "Missing serial number", "Completeness", "High")
    for i in m.sample(df, 0.03, col="serial_number"):
        m.set(df, i, "serial_number", df.at[i, "serial_number"].lower() + " ", t, "equipment_id",
              "Serial number lowercase/trailing space", "Consistency", "Low")
    for i in m.sample(df, 0.02, col="service_meter_hours", kind=int):
        m.set(df, i, "service_meter_hours", -int(df.at[i, "service_meter_hours"]) - 1, t, "equipment_id",
              "Negative service meter hours", "Validity", "Medium")
    for i in m.sample(df, 0.02):
        m.set(df, i, "model_year", rng.choice([2031, 1899, 0]), t, "equipment_id",
              "Impossible model year", "Validity", "Medium")
    for i in m.sample(df, 0.05):
        orig = df.at[i, "equipment_model"]
        m.set(df, i, "equipment_model", orig.replace("CAT ", "").lower(), t, "equipment_id",
              "Model text does not match product master", "Consistency", "Medium")
    for i in m.sample(df, 0.02):
        dup = m.duplicate(df, i, t, "equipment_id",
                          {"equipment_id": lambda x: "E9" + x[2:]}, "Same serial loaded twice", "High")
        df.loc[len(df)] = dup

    # --- Part -------------------------------------------------------------
    t, df = "part_master", raw["part_master"]
    for i in m.sample(df, 0.08):
        orig = df.at[i, "part_number"]
        m.set(df, i, "part_number", rng.choice([orig.replace("-", ""), orig.lower(), f" {orig}",
                                                 orig.replace("-", " "), "0" + orig]),
              t, "part_id", "Part number format not standardized", "Validity", "Medium")
    for i in m.sample(df, 0.06):
        m.set(df, i, "part_description", df.at[i, "part_description"].title().replace(" ", "  ", 1)
              if rng.random() < 0.5 else df.at[i, "part_description"].replace("FILTER", "FLTR").replace("ASSEMBLY", "ASSY"),
              t, "part_id", "Description abbreviations/spacing", "Consistency", "Low")
    for i in m.sample(df, 0.04):
        m.set(df, i, "part_category", None, t, "part_id", "Missing part category", "Completeness", "High")
    for i in m.sample(df, 0.04):
        m.set(df, i, "criticality", None, t, "part_id", "Missing criticality", "Completeness", "High")
    for i in m.sample(df, 0.03, col="unit_cost_usd", kind=float):
        m.set(df, i, "unit_cost_usd", rng.choice([0, -1 * df.at[i, "unit_cost_usd"], None]), t, "part_id",
              "Zero, negative, or missing unit cost", "Validity", "High")
    for i in m.sample(df, 0.02, col="unit_cost_usd", kind=float):
        m.set(df, i, "unit_cost_usd", f"${df.at[i, 'unit_cost_usd']:,.2f}", t, "part_id",
              "Unit cost stored as currency text", "Validity", "Medium")
    for i in m.sample(df, 0.03):
        m.set(df, i, "primary_supplier_id", rng.choice(["S099", None, "CAT"]), t, "part_id",
              "Invalid supplier reference", "Integrity", "High")
    for i in m.sample(df, 0.03):
        m.set(df, i, "unit_of_measure", rng.choice(["each", "Ea.", "PC", "ea"]), t, "part_id",
              "UOM not standardized", "Consistency", "Low")
    for i in m.sample(df, 0.04, col="created_date"):
        m.set(df, i, "created_date", messy_date(df.at[i, "created_date"], rng), t, "part_id",
              "Non-ISO date format", "Validity", "Low")
    for i in m.sample(df, 0.04):
        dup = m.duplicate(df, i, t, "part_id",
                          {"part_id": lambda x: "PT9" + x[3:],
                           "part_number": lambda x: x.replace("-", "")},
                          "Duplicate part (number entered without dash)", "High")
        df.loc[len(df)] = dup
    # Supersession: old number still active alongside its replacement (very common with Cat parts)
    for i in m.sample(df, 0.02):
        dup = m.duplicate(df, i, t, "part_id",
                          {"part_id": lambda x: "PT8" + x[3:],
                           "part_number": lambda x: x[:-1] + str((int(x[-1]) + 1) % 10) if x[-1].isdigit() else x,
                           "part_description": lambda x: x + " (SUPERSEDED)"},
                          "Superseded part number still active", "Medium")
        df.loc[len(df)] = dup

    # Shuffle rows so defects aren't clustered at the end
    for k in raw:
        raw[k] = raw[k].sample(frac=1, random_state=rng.randint(0, 10_000)).reset_index(drop=True)

    return raw, pd.DataFrame(m.log)


# ---------------------------------------------------------------------------

def write_excel(df: pd.DataFrame, path: Path, sheet: str) -> None:
    """One workbook per table: bold header, frozen top row, filters, sized columns."""
    from openpyxl.styles import Font, PatternFill

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name=sheet[:31], index=False)
        ws = xw.sheets[sheet[:31]]
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_cells in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[:500])
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 60)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--scale", type=float, default=1.0,
                        help="volume multiplier for equipment units and parts (e.g. 0.2 small, 3 large)")
    args = parser.parse_args()

    rng = random.Random(args.seed)
    golden = build_golden(rng, args.scale)
    raw, log = build_raw(golden, rng)

    for sub, data in (("golden", golden), ("raw", raw)):
        (OUT_DIR / sub).mkdir(parents=True, exist_ok=True)
        for name, df in data.items():
            df.to_csv(OUT_DIR / sub / f"{name}.csv", index=False)
            write_excel(df, OUT_DIR / sub / f"{name}.xlsx", name)
    log.to_csv(OUT_DIR / "dq_issue_log.csv", index=False)
    write_excel(log, OUT_DIR / "dq_issue_log.xlsx", "dq_issue_log")

    print(f"{'table':<26}{'golden':>8}{'raw':>8}{'issues':>8}")
    for name in golden:
        print(f"{name:<26}{len(golden[name]):>8}{len(raw[name]):>8}{(log.table == name).sum():>8}")
    print(f"\nTotal logged issues: {len(log)}  ->  {OUT_DIR}")


if __name__ == "__main__":
    main()
